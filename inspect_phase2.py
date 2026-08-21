import json
import docx
import os

p_nb3 = r"C:\Users\HP\Downloads\03_D1NAMO_Phase2_NonInvasive.ipynb"
if os.path.exists(p_nb3):
    with open(p_nb3, 'r', encoding='utf-8', errors='ignore') as f:
        nb3 = json.load(f)
    print(f"=== 03_D1NAMO_Phase2_NonInvasive.ipynb ({len(nb3.get('cells', []))} cells) ===")
    for idx, cell in enumerate(nb3.get('cells', [])):
        ct = cell.get('cell_type')
        src = "".join(cell.get('source', []))
        if ct == 'markdown':
            heads = [line for line in src.split('\n') if line.strip().startswith('#')]
            if heads:
                print(f"  [Cell {idx} MD]: {', '.join(heads)}")
        elif ct == 'code':
            first_line = src.strip().split('\n')[0] if src.strip() else ""
            print(f"  [Cell {idx} CODE]: {first_line[:80]}...")

p_doc2 = r"C:\Users\HP\Downloads\D1NAMO_Phase2_Report.docx"
if os.path.exists(p_doc2):
    doc2 = docx.Document(p_doc2)
    print(f"\n=== D1NAMO_Phase2_Report.docx ({len(doc2.paragraphs)} paras) ===")
    for p in doc2.paragraphs[:20]:
        if p.text.strip():
            print("  ", p.text.strip())
