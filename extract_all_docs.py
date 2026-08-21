import os
import sys

def read_pdf(pdf_path):
    import pypdf
    reader = pypdf.PdfReader(pdf_path)
    text = []
    for idx, page in enumerate(reader.pages):
        text.append(f"--- Page {idx+1} ---")
        text.append(page.extract_text() or "")
    return "\n".join(text)

def read_pptx(pptx_path):
    from pptx import Presentation
    prs = Presentation(pptx_path)
    text = []
    for idx, slide in enumerate(prs.slides):
        text.append(f"--- Slide {idx+1} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.text.strip():
                        text.append(paragraph.text.strip())
            elif shape.has_table:
                for row in shape.table.rows:
                    row_txt = " | ".join([cell.text.strip() for cell in row.cells])
                    text.append(row_txt)
    return "\n".join(text)

def read_docx(docx_path):
    import docx
    doc = docx.Document(docx_path)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_txt = " | ".join([cell.text.strip() for cell in row.cells])
            text.append(row_txt)
    return "\n".join(text)

def main():
    docs = [
        ("Phase 1 - DETAIL REPORT (1).pdf", read_pdf),
        ("Phase 1 - Glucose Forecasting (1).pdf", read_pdf),
        ("Wearable_AI_Diabetes_IR_Presentation (3).pptx", read_pptx),
        ("Wearable_AI_Learning_Guide (2).docx", read_docx),
    ]
    
    for filename, reader_fn in docs:
        path = os.path.join(r"d:\BTP", filename)
        out_path = os.path.join(r"d:\BTP", f"{os.path.splitext(filename)[0]}_extracted.txt")
        print(f"Reading {filename}...")
        try:
            content = reader_fn(path)
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(content)
            print(f"Saved {out_path} ({len(content)} chars)")
        except Exception as e:
            print(f"Error reading {filename}: {e}")

if __name__ == '__main__':
    main()
