import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def markdown_to_docx(md_path, docx_path, title_text):
    doc = Document()
    
    # Page Margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 51, 51)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        rows = [re.split(r'\s*\|\s*', l.strip().strip('|')) for l in t_lines if not re.match(r'^\s*\|?\s*[-:]+[-| :]*$', l)]
        if not rows or len(rows) < 1:
            return
        
        num_cols = max(len(r) for r in rows)
        tbl = doc.add_table(rows=len(rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, r in enumerate(rows):
            for c_idx in range(num_cols):
                cell_text = r[c_idx] if c_idx < len(r) else ""
                cell = tbl.cell(r_idx, c_idx)
                cell.text = cell_text
                
                # Format
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9.5)
                    if r_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        run.font.color.rgb = RGBColor(40, 40, 40)
                
                if r_idx == 0:
                    set_cell_background(cell, "1F4E79") # Deep Navy
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F2F5F8") # Light Slate
                else:
                    set_cell_background(cell, "FFFFFF")
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    for line in lines:
        stripped = line.rstrip()

        # Code block toggle
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph('\n'.join(code_lines))
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                for run in p.runs:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(30, 30, 30)
                code_lines = []
            else:
                if in_table:
                    flush_table(table_lines)
                    in_table = False
                    table_lines = []
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(stripped)
            continue

        # Tables
        if '|' in stripped and not stripped.startswith('#'):
            in_table = True
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                flush_table(table_lines)
                in_table = False
                table_lines = []

        # Headers
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(46, 117, 182)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            run = h.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(68, 84, 106)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            clean_text = stripped[2:]
            # handle simple bold
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            p.paragraph_format.space_after = Pt(2)
        elif stripped.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
        elif stripped:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            p.paragraph_format.space_after = Pt(4)

    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"Exported: {docx_path} ({os.path.getsize(docx_path)/(1024):.1f} KB)")

if __name__ == '__main__':
    markdown_to_docx(
        '01_NHANES_Strategic_Pivot_and_Plan_of_Action.md',
        '01_NHANES_Strategic_Pivot_and_Plan_of_Action.docx',
        'Strategic Pivot & Plan of Action'
    )
    markdown_to_docx(
        '02_NHANES_Comprehensive_EDA_and_Feature_Report.md',
        '02_NHANES_Comprehensive_EDA_and_Feature_Report.docx',
        'Comprehensive EDA Report'
    )
