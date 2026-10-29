import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_certificate_docx(project_title="Wearable AI for Insulin Resistance Screening"):
    doc = Document()

    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Add page border in oxml
        sectPr = section._sectPr
        pgBorders = parse_xml(
            f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
            f'<w:top w:val="double" w:sz="12" w:space="24" w:color="333333"/>'
            f'<w:left w:val="double" w:sz="12" w:space="24" w:color="333333"/>'
            f'<w:bottom w:val="double" w:sz="12" w:space="24" w:color="333333"/>'
            f'<w:right w:val="double" w:sz="12" w:space="24" w:color="333333"/>'
            f'</w:pgBorders>'
        )
        sectPr.append(pgBorders)

    # Header: Indian Institute of Technology Guwahati
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(12)
    p_header.paragraph_format.space_after = Pt(18)
    run_header = p_header.add_run("Indian Institute of Technology Guwahati")
    run_header.font.name = "Times New Roman"
    run_header.font.size = Pt(16)
    run_header.bold = True
    run_header.font.color.rgb = RGBColor(0, 0, 0)

    # Logo
    logo_path = r"d:\BTP\iitg_logo.png"
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(28)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(1.5))

    # Title: Certificate
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(28)
    run_title = p_title.add_run("Certificate")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(22)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0)

    # Main Body Paragraph 1
    p_body1 = doc.add_paragraph()
    p_body1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body1.paragraph_format.line_spacing = 1.35
    p_body1.paragraph_format.space_after = Pt(14)

    run = p_body1.add_run("This is to certify that ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    run_name = p_body1.add_run("Mr. Sarvagya Tiwari (Roll No. 230122047)")
    run_name.font.name = "Times New Roman"
    run_name.font.size = Pt(12)
    run_name.bold = True

    run = p_body1.add_run(" has been working under my supervision as a Project Student in my research group since July 2025.")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    # Main Body Paragraph 2
    p_body2 = doc.add_paragraph()
    p_body2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body2.paragraph_format.line_spacing = 1.35
    p_body2.paragraph_format.space_after = Pt(14)

    run = p_body2.add_run("His project, entitled ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    run_proj = p_body2.add_run(f"“{project_title}”")
    run_proj.font.name = "Times New Roman"
    run_proj.font.size = Pt(12)
    run_proj.bold = True

    run = p_body2.add_run(", is being forwarded for submission. I hereby certify that he has satisfactorily completed all the requirements in accordance with the rules and regulations of the Institute pertaining to the conduct of such investigations.")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    # Main Body Paragraph 3
    p_body3 = doc.add_paragraph()
    p_body3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body3.paragraph_format.line_spacing = 1.35
    p_body3.paragraph_format.space_after = Pt(40)

    run = p_body3.add_run("To the best of my knowledge, the results presented in this project report have not been submitted elsewhere for the award of any degree.")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    # Supervisor Info & Signature
    p_prof = doc.add_paragraph()
    p_prof.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_prof.paragraph_format.line_spacing = 1.15
    p_prof.paragraph_format.space_after = Pt(2)

    r_prof = p_prof.add_run("Prof. Ritika Jain")
    r_prof.font.name = "Times New Roman"
    r_prof.font.size = Pt(12)
    r_prof.bold = True

    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.line_spacing = 1.15
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("Indian Institute of Technology Guwahati\ne-mail: rjain1992@iitg.ac.in")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(11)

    p_date = doc.add_paragraph()
    p_date.paragraph_format.line_spacing = 1.15
    p_date.paragraph_format.space_after = Pt(30)
    r_date = p_date.add_run("Date: May 2026")
    r_date.font.name = "Times New Roman"
    r_date.font.size = Pt(11)

    # Signature line
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(20)
    p_sig.paragraph_format.space_after = Pt(2)
    r_line = p_sig.add_run("_____________________________")
    r_line.font.name = "Times New Roman"
    r_line.font.size = Pt(12)

    p_sig_txt = doc.add_paragraph()
    p_sig_txt.paragraph_format.space_after = Pt(0)
    r_sig = p_sig_txt.add_run("Signature of Supervisor")
    r_sig.font.name = "Times New Roman"
    r_sig.font.size = Pt(12)
    r_sig.bold = True

    out_path = r"d:\BTP\IITG_BTP_Certificate_Sarvagya_Tiwari.docx"
    try:
        doc.save(out_path)
        print(f"Generated DOCX: {out_path} with title: {project_title}")
    except PermissionError:
        out_path_v2 = r"d:\BTP\IITG_BTP_Certificate_ShortTitle.docx"
        doc.save(out_path_v2)
        print(f"File was locked by Word. Saved to: {out_path_v2} with title: {project_title}")

if __name__ == "__main__":
    create_certificate_docx("Wearable AI for Insulin Resistance Screening")
